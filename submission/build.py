"""The idea deck: PS 26067's submitted deck (72/90) with its text and pictures replaced.

    python submission/figures.py     # pictures first
    python submission/build.py       # -> submission/final/Satelight-SIH2026.pptx
    powershell -ExecutionPolicy Bypass -File submission/topdf.ps1

The layout, fonts, colours, masthead and footer are the P3 deck's, unchanged: its final
.pptx is copied in as submission/base-p3.pptx (not committed; from
SIH26P3/submission/sih/final/VVater-SIH2026-final.pptx). Every shape keeps its place and
its run formatting; only the words and the pictures change. Every number is read from
data/eval-latest.json, never typed (hard rule 1), and every claim about a score is
asserted against it before the deck is written.
"""

import copy
import json
import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
BASE = HERE / "base-p3.pptx"
P3 = Path(r"C:\Users\Cleo\Desktop\SIH26P3\submission\sih\final\VVater-SIH2026-final.pptx")
FIG = HERE / "figures"
OUT = HERE / "final" / "Satelight-SIH2026.pptx"
REPO = "https://github.com/Kukyos/Satelight"

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

# ------------------------------------------------------------------ numbers, from the harness

E = json.loads((ROOT / "data" / "eval-latest.json").read_text().replace("NaN", "null"))
HEAD = f"{E['selection']['headline']} (selected)"
C = E["argo_counts"]
DEPTHS = E["depths"]
BOX, BOB, AS = "North Indian Ocean (whole box)", "Bay of Bengal", "Arabian Sea"


def rmse(region, name, depth):
    return E["argo"][region][name][DEPTHS.index(depth)]["rmse"]


def k(x: int) -> str:
    return f"{x:,}"


# Claims the slides make in words, checked here so a new harness run cannot silently
# falsify them.
for region in (BOX, BOB, AS):
    for d in DEPTHS:
        if d <= 300:
            assert rmse(region, HEAD, d) < rmse(region, "climatology", d), (region, d, "floor")
            assert rmse(region, HEAD, d) > rmse(region, "GLORYS (ceiling)", d), (region, d, "ceiling")
for d in (75, 100, 125, 150):
    assert E["argo"][BOX][HEAD][DEPTHS.index(d)]["bias"] > 0, "thermocline runs warm"

TCHP = E["hazard"]["TCHP (kJ/cm²)"]
tchp = {r: {n: TCHP[r][n]["rmse"] for n in ("climatology", HEAD, "GLORYS (ceiling)")} for r in (BOX, BOB, AS)}
for r in tchp:
    assert tchp[r][HEAD] < tchp[r]["climatology"], r
MOCHA, BIP = E["cyclones"]["Mocha"], E["cyclones"]["Biparjoy"]
m_t, b_t = MOCHA["tchp_box_mean"], BIP["tchp_box_mean"]
assert m_t["reconstruction"]["after"] < m_t["reconstruction"]["before"]
assert b_t["reconstruction"]["after"] < b_t["reconstruction"]["before"]
# The float in the slide-2 crop, as capture.cjs picked it; its paperwork goes in the caption.
CAST = json.loads((FIG / "shot-profile.json").read_text())
GAP = E["gap_closed"][BOX][HEAD]
gap0 = GAP[0]
EMB = E["embedding"]
S4_PROBE_RUN = E["embedding_screen"]["chosen"]
_probe = EMB[S4_PROBE_RUN]["mld_probe"]
S4_PROBE = {"embedding": next(v["r2_test"] for kk, v in _probe.items() if kk.startswith("embedding")),
            "raw": next(v["r2_test"] for kk, v in _probe.items() if kk.startswith("raw"))}
# The slide says "half met": each S4 check is met by some embedding, no one embedding meets both.
_verdicts = [v["verdict"] for v in EMB.values() if v.get("verdict")]
assert not any(all(v.values()) for v in _verdicts), "S4 is met in full: reword slide 4"
assert all(any(v[c] for v in _verdicts) for c in _verdicts[0]), "a check is met by no run"
assert S4_PROBE["embedding"] > S4_PROBE["raw"]


def f3(x):
    return f"{x:.3f}"


def f1(x):
    return f"{x:.1f}"


def f2(x):
    return f"{x:.2f}"


# ------------------------------------------------------------------ the words

TITLE_PS = ("OceanEmbed - Satellite Embedding-Based Deep Learning Framework for Reconstruction "
            "of Subsurface Ocean Temperature from Surface Satellite Observations")

S1 = {"TextBox 11": [
    "Problem Statement ID - SIH26066", "", f"Problem Statement Title - {TITLE_PS}", "",
    "Theme - Disaster Management", "", "PS Category - Software", "", "Team ID - K26125", "",
    "Team Name - Team Null"]}

S2 = {
    "Group 4/TextBox 6": ["Satelight: Ocean from Space"],
    "TextBox 78": [["Eight satellite surface fields in, ",
                    "temperature at 15 depths to 1,000 m out, every day at 0.25°",
                    ", over the North Indian Ocean, through a learned embedding."]],
    "TextBox 79": [["Working prototype (idea stage), captured.",
                    " The Bay of Bengal on 2023-05-17, after Cyclone Mocha, reconstructed from "
                    f"satellites alone to 1,000 m. Bottom right: held-out Argo float {CAST['platform']} "
                    f"cycle {CAST['cycle']}, data mode {CAST['data_mode']}, {CAST['levels_rejected']} "
                    f"levels failed QC, file {CAST['source_file'].split('/')[-1].split(' ')[0]}."]],
    "TextBox 81": [["Inputs:", " SST, salinity, sea level (SSH, SLA), currents and winds, "
                    "harmonised onto one 0.25° daily grid"]],
    "TextBox 82": [["Model:", " CNN + transformer encoder makes the embedding; a decoder "
                    "turns it into 15 levels; trained on GLORYS12, 2010–2020"]],
    "TextBox 83": [["Proof:", f" {k(C['casts_used'])} held-out Argo casts, 2023–24, scored "
                    "per depth and basin beside climatology and GLORYS"]],
    "TextBox 85": ["The expected solution, point by point, and what is built."],
    "TextBox 86": ["Preprocessing and harmonisation of multi-source data"],
    "TextBox 87": ["one fetch per source, block means onto 0.25° daily; every assumption "
                   "travels in the file; a re-fetch is bit-identical"],
    "TextBox 88": ["A satellite embedding engine"],
    "TextBox 89": ["each cell's surface state compressed to 32 numbers, saved per day as a "
                   "file you can open, cluster and probe"],
    "TextBox 90": ["Deep-learning reconstruction at 15 standard depths"],
    "TextBox 91": ["exactly 0 to 1,000 m on the asked levels; cells shallower than a level "
                   "are masked, never extrapolated"],
    "TextBox 92": ["Validation with independent Argo"],
    "TextBox 93": [f"RMSE, bias and r per depth and basin; {k(C['levels_rejected_by_qc'])} "
                   "levels failing QC counted, never dropped"],
    "TextBox 94": ["Working PoC over the Bay of Bengal and Arabian Sea"],
    "TextBox 95": ["a 3D viewer: the inputs, the embedding, the reconstruction beside GLORYS, "
                   "and the held-out floats, day by day"],
    "TextBox 96": ["Also asked for by name, and built"],
    "TextBox 97": ["Daily NetCDF output, 0.25° · INCOIS gridded Argo as a second "
                   "reference · a run without currents and winds · every run from a config "
                   "and a seed"],
    "TextBox 99": ["Scored between a floor and a ceiling"],
    "TextBox 100": ["every number beside climatology and GLORYS; at the surface it closes "
                    f"{gap0:.0%} of the gap between them"],
    "TextBox 101": ["What a cyclone feeds on"],
    "TextBox 48": ["20 °C and 26 °C isotherm depths and heat potential, "
                   "each checked against Argo"],
    "TextBox 102": ["Sees a cyclone's cold wake"],
    "TextBox 103": ["Mocha and Biparjoy, 2023, reconstructed from satellites only"],
    "TextBox 104": ["The reference is checked too"],
    "TextBox 105": ["INCOIS gridded Argo range-tested and scored against Argo itself, "
                    "before anyone is scored against it"],
    "TextBox 106": ["No leakage, by design"],
    "TextBox 107": ["contiguous time blocks; GLORYS never an input; the test years never "
                    "touched in training"],
    "TextBox 110": ["An embedding you can open"],
    "TextBox 111": ["saved per day, clustered and probed; where it falls short, the deck says so"],
    "TextBox 108": ["a float's profile vs every model"],
    "TextBox 109": ["Mocha's cold wake, satellites only"],
}

S3 = {
    "TextBox 88": ["Software architecture: three layers, end to end"],
    "TextBox 90": ["forecaster, researcher, fisheries or disaster officer"],
    "TextBox 91": ["Picks a basin, a day", "2023–24, held out", " ",
                   "Switches the field", "ours, GLORYS, error", " ",
                   "Clicks a float", "vs every model", " ",
                   "Takes the day away", "NetCDF, 15 levels"],
    "TextBox 92": ["1. Presentation layer: viewer"],
    "TextBox 93": ["runs in the browser, served by the API"],
    "TextBox 94": ["Tech: TypeScript, Vite"],
    "TextBox 95": ["Satellite inputs and embedding"],
    "TextBox 96": ["the eight fields and the embedding, as maps"],
    "TextBox 97": ["3D temperature block on the globe"],
    "TextBox 98": ["15 levels to 1,000 m; floats as upright lines"],
    "TextBox 99": ["Tech: CesiumJS, WebGL"],
    "TextBox 100": ["Float profile chart"],
    "TextBox 101": ["a held-out Argo cast next to every model"],
    "TextBox 102": ["Tech: SVG"],
    "TextBox 103": ["Skill by depth"],
    "TextBox 104": ["RMSE against Argo per basin, from the harness"],
    "TextBox 105": ["Tech: SVG"],
    "TextBox 106": ["Error view"],
    "TextBox 107": ["reconstruction minus GLORYS in 3D, on a diverging scale"],
    "TextBox 108": ["2. Application layer: model"],
    "TextBox 109": ["Python, runs on one 8 GB GPU"],
    "TextBox 110": ["Tech: PyTorch, FastAPI"],
    "TextBox 111": ["Step 1: fetch and harmonise"],
    "TextBox 112": ["eight inputs and the target onto one 0.25° daily grid"],
    "TextBox 113": ["Tech: copernicusmarine, earthaccess, xarray"],
    "TextBox 114": ["Step 2: encode the surface"],
    "TextBox 115": ["CNN + transformer: 32 numbers per cell per day"],
    "TextBox 116": ["Tech: PyTorch"],
    "TextBox 117": ["Step 3: decode 15 depths"],
    "TextBox 118": ["0 to 1,000 m; nothing below the sea floor"],
    "TextBox 119": ["Tech: PyTorch"],
    "TextBox 120": ["Step 4: score on held-out Argo"],
    "TextBox 121": ["QC flags kept; per depth and basin; heat potential"],
    "TextBox 122": ["Tech: NumPy, GSW (TEOS-10)"],
    "TextBox 123": ["Output writer"],
    "TextBox 124": ["daily NetCDF and embedding files, provenance attached"],
    "TextBox 125": ["Tech: xarray, NetCDF4"],
    "TextBox 126": ["3. Data layer (APIs)"],
    "TextBox 127": ["public sources, read over the internet"],
    "TextBox 128": ["Copernicus Marine", "SST, salinity, sea level; GLORYS12"],
    "TextBox 129": ["NASA PO.DAAC", "OSCAR currents, CCMP winds"],
    "TextBox 130": ["Ifremer", "Argo floats, held out"],
    "TextBox 131": ["INCOIS ERDDAP", "gridded Argo, 1°"],
    "TextBox 132": ["IMD", "cyclone dates, Mocha, Biparjoy"],
    "TextBox 133": ["Open-source stack", "PyTorch, xarray, CesiumJS", "free, no licence cost"],
    "TextBox 134": ["3D data"],
    "TextBox 135": ["One command starts it: start.bat or start.sh.",
                    "Runs on one 8 GB GPU"],
    "Group 79/TextBox 81": [["What the model writes for each cell and day:", "  temperature at 0, 5, 10, "
                             "20, 30, 50, 75, 100, 125, 150, 200, 300, 500, 700 and 1,000 m"]],
    "Group 64/TextBox 66": ["uses"],
    "Group 68/TextBox 70": ["request"],
    "Group 72/TextBox 74": ["subset only"],
    "Group 76/TextBox 78": ["downloads"],
}

S4 = {
    "TextBox 14": ["ANALYSIS OF FEASIBILITY: MEASURED BY THE EVAL HARNESS"],
    "TextBox 17": [k(C["casts_used"])],
    "TextBox 18": ["held-out Argo casts scored, 2023–24; the model never saw one of them"],
    "TextBox 21": [str(E["test_days"])],
    "TextBox 22": ["test days reconstructed and scored, one file per day"],
    "TextBox 25": [k(C["levels_rejected_by_qc"])],
    "TextBox 26": ["Argo levels failing QC: counted in the report, never silently dropped"],
    "TextBox 29": [f"{f1(tchp[BOX]['climatology'])} → {f1(tchp[BOX][HEAD])}"],
    "TextBox 30": ["kJ/cm² RMSE in cyclone heat potential vs Argo: climatology → ours "
                   f"(GLORYS {f1(tchp[BOX]['GLORYS (ceiling)'])})"],
    "TextBox 31": ["WHAT THE HARNESS FOUND: ERROR AGAINST HELD-OUT ARGO, BY DEPTH"],
    "TextBox 66": [["From the surface to 300 m Satelight beats climatology in both basins and "
                    "stays behind GLORYS, which assimilated these floats. ",
                    "The thermocline (75–150 m) is the hard part",
                    ": errors peak there and run warm."]],
    "TextBox 67": ["CHALLENGES AND RISKS, AND OUR STRATEGIES"],
    "TextBox 76": ["GLORYS already contains the Argo floats"],
    "TextBox 79": ["GLORYS is never an input; it is scored as the ceiling, not a rival."],
    "TextBox 82": ["Few years of satellite salinity"],
    "TextBox 85": ["Contiguous time blocks; recipe chosen on the validation years only."],
    "TextBox 88": ["The test years are warmer"],
    "TextBox 91": ["An anomaly target; the warm bias at 75–150 m is reported, not hidden."],
    "TextBox 94": ["Sources on different grids and clocks"],
    "TextBox 97": ["Block means only, never upsampled; each assumption written into the file."],
    "TextBox 100": ["The embedding proof is half met"],
    "TextBox 103": [f"{S4_PROBE_RUN}'s embedding beats the raw inputs on mixed-layer depth "
                    f"(R² {f3(S4_PROBE['embedding'])} vs {f3(S4_PROBE['raw'])}); a second encoder is next."],
    "TextBox 112": ["SUSTAINABILITY", "0 Rs running cost; free public data",
                    "New source = one fetch function", "Every run from a config and a seed",
                    "Open source; issues and PRs open"],
    "TextBox 112#2": ["ROADMAP"],
    "Pentagon 113": ["NOW", "prototype, 2 basins"],
    "Pentagon 114": ["3 MO", "embedding proof, salinity"],
    "Pentagon 115": ["6 MO", "daily feed for INCOIS"],
    "Pentagon 116": ["12 MO", "whole Indian Ocean"],
}

S5 = {
    "TextBox 14": ["EACH TYPE OF USER AND THE IMPACT ON THEM"],
    "TextBox 15": ["INCOIS cyclone forecasters"],
    "TextBox 26": ["Satelight"],
    "TextBox 37": ["GLORYS"],
    "TextBox 38": [f"Cyclone Mocha, {MOCHA['before']} to {MOCHA['after'][-2:]}: box-mean heat "
                   f"potential {f1(m_t['reconstruction']['before'])} → "
                   f"{f1(m_t['reconstruction']['after'])} kJ/cm², satellites only."],
    "TextBox 39": ["Fisheries"],
    "TextBox 50": ["Where the warm layer ends, every day at 0.25°: the thermocline depth "
                   "that pelagic fish follow, between the floats."],
    "TextBox 51": ["Researchers and modellers"],
    "TextBox 62": ["Reconstruction minus GLORYS in 3D, day by day: where satellites can see "
                   "below the surface and where they cannot."],
    "TextBox 63": ["Public and policymakers"],
    "TextBox 74": ["The embedding as colour: each day's ocean summarised by the model, "
                   "water that looks alike painted alike."],
    "TextBox 75": ["BENEFITS"],
    "TextBox 81": ["Disaster preparedness: a cyclone's fuel, daily, between the floats"],
    "TextBox 82": ["Cyclone wakes seen below the surface from satellites alone"],
    "TextBox 83": ["Open data and code for Indian ocean science and teaching"],
    "TextBox 89": ["Uses satellites already in orbit: no new instruments to buy"],
    "TextBox 90": ["One 8 GB GPU; free public data; open source"],
    "TextBox 91": ["Fills the gaps between Argo floats, which are costly to deploy"],
    "TextBox 97": ["Upper-ocean heat content every day, for climate monitoring"],
    "TextBox 98": ["Marine heatwaves below the surface, not only at it"],
    "TextBox 99": ["Every score shown beside climatology and GLORYS, so decisions know the limits"],
    "TextBox 99#2": ["Scale: one model covers the whole North Indian Ocean box; a new region is "
                     "a config line and a retrain. Next: salinity at depth, heatwaves."],
}

REFS = [  # (number, text, grey tail, link)
    ("Copernicus Marine GLORYS12 reanalysis, daily. ", "The training target and the ceiling.",
     "data.marine.copernicus.eu/product/GLOBAL_MULTIYEAR_PHY_001_030"),
    ("OSTIA SST and multi-observation SSS, reprocessed. ", "Two inputs.",
     "data.marine.copernicus.eu"),
    ("DUACS sea level, SLA and ADT. ", "Input; the common grid.",
     "data.marine.copernicus.eu/product/SEALEVEL_GLO_PHY_L4_MY_008_047"),
    ("OSCAR L4 v2.0 surface currents, NASA PO.DAAC. ", "Input.",
     "podaac.jpl.nasa.gov/dataset/OSCAR_L4_OC_FINAL_V2.0"),
    ("CCMP v3.1 10 m winds, Remote Sensing Systems. ", "Input.",
     "podaac.jpl.nasa.gov/dataset/CCMP_WINDS_10M6HR_L4_V3.1"),
    ("Argo profiles via Ifremer ERDDAP. ", "The held-out validation.",
     "erddap.ifremer.fr/erddap"),
    ("INCOIS ERDDAP, Argo 10-day analysis. ", "Second reference.",
     "erddap.incois.gov.in/erddap"),
    ("IMD RSMC New Delhi, Mocha and Biparjoy, 2023. ", "Storm dates.",
     "rsmcnewdelhi.imd.gov.in"),
    ("Wong et al., Argo Quality Control Manual. ", "QC flags, range test.",
     "doi.org/10.13155/33951"),
    ("Ronneberger et al. (2015), U-Net. ", "The convolutional encoder.",
     "doi.org/10.1007/978-3-319-24574-4_28"),
    ("Vaswani et al. (2017), Attention. ", "The hybrid's transformer.",
     "arxiv.org/abs/1706.03762"),
    ("CF Conventions for NetCDF. ", "The output file.", "cfconventions.org"),
    ("IOC, SCOR, IAPSO (2010), TEOS-10. ", "Depth, heat capacity.", "teos-10.org"),
    ("Leipper & Volgenau (1972), J. Phys. Oceanogr. 2. ", "Heat above 26 °C.",
     "doi.org/10.1175/1520-0485(1972)002<0218:HHPOTG>2.0.CO;2"),
    ("Thyng et al. (2016), Oceanography 29(3). ", "cmocean palettes.",
     "doi.org/10.5670/oceanog.2016.66"),
    ("CesiumJS. ", "The globe and renderer.", "cesium.com/platform/cesiumjs"),
    ("Gray (1968), Mon. Weather Rev. 96, 669–700. ", "Cyclones need water above ~26 °C.", ""),
    ("Price (1981), J. Phys. Oceanogr. 11, 153–175. ", "A cyclone cools the sea.", ""),
    ("Schott & McCreary (2001), Prog. Oceanogr. 51, 1–123. ", "The monsoon current reverses.", ""),
    ("Loshchilov & Hutter (2019), Decoupled weight decay. ", "The optimiser.",
     "arxiv.org/abs/1711.05101"),
    ("Lellouche et al. (2021), Front. Earth Sci. 9. ", "GLORYS12 and what it assimilates.",
     "doi.org/10.3389/feart.2021.698876"),
    ("Argo Program, \"How Argo floats work\". ", "Floats.", "argo.ucsd.edu"),
    ("SIH 2026 problem statement 26066, quoted verbatim. ", "",
     "github.com/Kukyos/Satelight/blob/main/docs/01-problem-statement.md"),
]

S6 = {
    "TextBox 20": [REPO.removeprefix("https://")],
    "TextBox 64": ["DATA AND METHODS"],
    "TextBox 105": ["WHAT THE SCIENCE CITES"],
}
# Reference rows: (number box, text box, link box) per slot, in P3's order.
REF_SLOTS = [(24, 25, 26), (29, 30, 31), (34, 35, 36), (39, 40, 41), (44, 45, 46), (49, 50, 51),
             (54, 55, 56), (59, 60, 61), (67, 68, 69), (72, 73, 74), (77, 78, 79), (82, 83, 84),
             (87, 88, 89), (92, 93, 94), (97, 98, 99), (102, 103, 104), (108, 109, None),
             (112, 113, None), (116, 117, None), (120, 121, None), (124, 125, None),
             (128, 129, 130), (133, 134, 135)]
for (n, t, l), (text, tail, link) in zip(REF_SLOTS, REFS):
    S6[f"TextBox {t}"] = [[text, tail]]
    if l:
        S6[f"TextBox {l}"] = [link]

SLIDES = [S1, S2, S3, S4, S5, S6]

# Pictures: slide index -> {shape path: figure}. The logo is one shared part.
PICTURES = {
    1: {"Group 14/Freeform 15": "deck-hero.jpg", "Group 55/Freeform 56": "deck-profile.jpg",
        "Group 65/Freeform 66": "deck-mocha-small.jpg"},
    4: {"Group 16/Freeform 17": "deck-mocha-rec.jpg", "Group 27/Freeform 28": "deck-mocha-glorys.jpg",
        "Group 40/Freeform 41": "deck-arabian.jpg", "Group 52/Freeform 53": "deck-error.jpg",
        "Group 64/Freeform 65": "deck-embedding.jpg"},
}
LOGO = {1: "Group 12/Freeform 13", 2: "Group 18/Freeform 19", 3: "Group 12/Freeform 13",
        4: "Group 12/Freeform 13", 5: "Group 15/Freeform 16"}

# Slide 4's bar chart was drawn from shapes; it is replaced by one picture.
CHART_SHAPES = {"AutoShape 32", "AutoShape 34", "AutoShape 36", "Group 38", "Group 41", "Group 45",
                "Group 48", "Group 52", "Group 55", "Group 62", "Group 64",
                *(f"TextBox {i}" for i in (33, 35, 37, 40, 43, 44, 47, 50, 51, 54, 57, 58, 59, 60, 61))}

# ------------------------------------------------------------------ mechanics


def shapes_by_path(slide) -> dict:
    """Every shape by its group path; a repeated name gets '#2', '#3'."""
    out = {}

    def walk(shs, prefix):
        for sh in shs:
            key = prefix + sh.name
            n = 2
            while key in out:
                key = f"{prefix}{sh.name}#{n}"
                n += 1
            out[key] = sh
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes, prefix + sh.name + "/")
    walk(slide.shapes, "")
    return out


def put(shape, paras) -> None:
    """Replace the text, keeping each paragraph's and each run's formatting. Paragraph j
    takes template paragraph min(j, last); segment i of it takes template run min(i, last)."""
    body = shape.text_frame._txBody
    tmpl = body.findall(f"{A}p")
    for p in tmpl:
        body.remove(p)
    for j, para in enumerate(paras):
        segs = [para] if isinstance(para, str) else para
        src = tmpl[min(j, len(tmpl) - 1)]
        runs = src.findall(f"{A}r")
        p = copy.deepcopy(src)
        for r in p.findall(f"{A}r") + p.findall(f"{A}br") + p.findall(f"{A}fld"):
            p.remove(r)
        end = p.find(f"{A}endParaRPr")
        for i, text in enumerate(segs):
            if not text:
                continue
            if runs:
                r = copy.deepcopy(runs[min(i, len(runs) - 1)])
            else:  # an empty template paragraph: style the run like its end mark
                r = etree.SubElement(p, f"{A}r")
                if end is not None:
                    rpr = copy.deepcopy(end)
                    rpr.tag = f"{A}rPr"
                    r.append(rpr)
                etree.SubElement(r, f"{A}t")
            r.find(f"{A}t").text = text
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        body.append(p)


def swap(slide, shape, image: Path) -> None:
    """Point a picture-filled shape at a new image, shown whole (the figure is already
    cropped to the frame's aspect)."""
    _, rid = slide.part.get_or_add_image_part(str(image))
    blip = shape._element.find(f".//{A}blip")
    blip.set(R_EMBED, rid)
    src = blip.getparent().find(f"{A}srcRect")
    if src is not None:
        blip.getparent().remove(src)


def build(out: Path = OUT) -> Path:
    if not BASE.exists():
        shutil.copy2(P3, BASE)
    prs = Presentation(str(BASE))
    for i, (slide, texts) in enumerate(zip(prs.slides, SLIDES)):
        shapes = shapes_by_path(slide)
        missing = [k for k in texts if k not in shapes]
        assert not missing, (i + 1, missing)
        for key, paras in texts.items():
            put(shapes[key], paras)
        for key, fig in PICTURES.get(i, {}).items():
            swap(slide, shapes[key], FIG / fig)
        if i in LOGO:
            swap(slide, shapes[LOGO[i]], FIG / "deck-logo.png")
        if i == 3:
            for key in CHART_SHAPES:
                el = shapes[key]._element
                el.getparent().remove(el)
            slide.shapes.add_picture(str(FIG / "deck-rmse.png"), Emu(502920), Emu(3590000),
                                     width=Emu(8947404))
        if i == 5:  # the reference numbers are P3's 1..23 already; nothing else to renumber
            pass
    out.parent.mkdir(exist_ok=True)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    import sys
    # `build.py <path>` writes elsewhere, for when the deck is open in PowerPoint.
    print("deck:", build(Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else OUT))
